"""
Convert a (password-protected) statement PDF into an Excel file.

Everything runs locally on your machine - nothing is uploaded anywhere.
The password is asked at runtime and is never saved or printed.

Usage:
    python app.py                                  (window with file picker)
    python statement_to_excel.py statement.pdf     (command line)
    python statement_to_excel.py statement.pdf -o output.xlsx

Requirements:
    pip install pdfplumber openpyxl
"""

import argparse
import getpass
import re
import sys
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import kbank

# A transaction line usually starts with a date like 01/09/2026, 01-09-26, 01 Sep 2026
DATE_RE = re.compile(
    r"^\s*(\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}|\d{1,2}\s+[A-Za-z]{3,9}\.?\s+\d{2,4})"
)
AMOUNT_RE = re.compile(r"-?\(?\d{1,3}(?:,\d{3})*(?:\.\d{2})\)?-?")


class PasswordError(Exception):
    """The PDF is locked and the password is missing or wrong."""


def _is_password_error(exc):
    """pdfplumber wraps pdfminer's password error, so walk the exception chain."""
    while exc is not None:
        if "Password" in type(exc).__name__ or "Password" in repr(exc):
            return True
        exc = exc.__cause__ or exc.__context__
    return False


def open_pdf(path, password=None):
    """Open the PDF. Raises PasswordError if it is locked and the password doesn't work."""
    try:
        return pdfplumber.open(path, password=password or "")
    except Exception as exc:
        if _is_password_error(exc):
            raise PasswordError("Wrong password" if password else "Password required") from exc
        raise


def is_locked(path):
    """True if the PDF needs a password to open."""
    try:
        open_pdf(path).close()
        return False
    except PasswordError:
        return True


def clean(cell):
    if cell is None:
        return ""
    return re.sub(r"\s+", " ", str(cell)).strip()


def to_number(text):
    """'1,234.50' -> 1234.5, '(50.00)' or '50.00-' -> -50.0; otherwise return text unchanged."""
    t = text.replace(",", "").strip()
    neg = (t.startswith("(") and t.endswith(")")) or t.endswith("-") or t.startswith("-")
    t = t.strip("()-")
    try:
        value = float(t)
    except ValueError:
        return text
    return -value if neg else value


def extract_page_tables(page, page_no):
    """Strategy 1: use the table structure found in the PDF."""
    rows = []
    for table in page.extract_tables():
        for row in table:
            cells = [clean(c) for c in row]
            if any(cells):
                rows.append([page_no] + cells)
    return rows


class LineParser:
    """Strategy 2: parse text lines that start with a date (for statements without table borders)."""

    def __init__(self):
        self.rows, self.raw = [], []
        self.current = None

    def feed_page(self, page, page_no):
        text = page.extract_text() or ""
        for line in text.splitlines():
            self.raw.append([page_no, line])
            m = DATE_RE.match(line)
            if m:
                date = m.group(1)
                rest = line[m.end():].strip()
                amounts = AMOUNT_RE.findall(rest)
                desc = rest
                for a in amounts:
                    desc = desc.replace(a, "", 1)
                self.current = [page_no, date, clean(desc)] + amounts
                self.rows.append(self.current)
            elif self.current is not None and line.strip() and not AMOUNT_RE.search(line):
                # Continuation of the previous description (multi-line description)
                self.current[2] = clean(self.current[2] + " " + line)
            else:
                self.current = None


def write_sheet(ws, header, rows):
    ws.append(header)
    for row in rows:
        ws.append([to_number(c) if isinstance(c, str) and AMOUNT_RE.fullmatch(c.strip()) else c
                   for c in row])
    # Styling
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    width_cols = max((len(r) for r in [header] + rows), default=0)
    for i in range(1, width_cols + 1):
        letter = get_column_letter(i)
        longest = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(longest + 2, 8), 60)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"


def convert(pdf_path, out_path, password=None, progress=None):
    """
    Convert pdf_path to an Excel file at out_path.
    progress(done_pages, total_pages) is called after each page, if given.
    Returns a dict with counts for the summary.
    """
    table_rows = []
    lines = LineParser()

    with open_pdf(pdf_path, password) as pdf:
        total = len(pdf.pages)
        result = kbank.parse(pdf, progress)
        if result is not None:
            wb = Workbook()
            kbank.write_workbook(wb, result)
            wb.save(out_path)
            checks = kbank.summarize(result)
            return {"format": "KBank", "pages": total,
                    "transactions": len(result["transactions"]),
                    "account": result["info"].get("account"),
                    "skipped": len(result["skipped"]),
                    "checks": [(label, ok) for label, _, _, _, ok in checks]}

        for page_no, page in enumerate(pdf.pages, 1):
            table_rows += extract_page_tables(page, page_no)
            lines.feed_page(page, page_no)
            if progress:
                progress(page_no, total)

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    max_amounts = max((len(r) - 3 for r in lines.rows), default=0)
    write_sheet(ws, ["Page", "Date", "Description"] +
                [f"Amount {i + 1}" for i in range(max_amounts)], lines.rows)

    if table_rows:
        ws_t = wb.create_sheet("Tables")
        n = max(len(r) for r in table_rows) - 1
        write_sheet(ws_t, ["Page"] + [f"Col {i + 1}" for i in range(n)], table_rows)

    write_sheet(wb.create_sheet("Raw text"), ["Page", "Line"], lines.raw)
    wb.save(out_path)

    return {"format": "generic", "pages": total, "transactions": len(lines.rows),
            "table_rows": len(table_rows), "text_lines": len(lines.raw)}


def ask_password():
    if sys.stdin.isatty():
        return getpass.getpass("PDF password (input hidden): ")
    return sys.stdin.readline().rstrip("\r\n")  # password piped in (e.g. from a script)


def main():
    ap = argparse.ArgumentParser(description="Convert a statement PDF to Excel.")
    ap.add_argument("pdf", type=Path)
    ap.add_argument("-o", "--output", type=Path)
    ap.add_argument("--debug", metavar="TEXT",
                    help="write how the lines around TEXT (e.g. a date 04-09-26) are read "
                         "to debug.txt, instead of converting")
    args = ap.parse_args()

    if not args.pdf.exists():
        sys.exit(f"File not found: {args.pdf}")
    out = args.output or args.pdf.with_suffix(".xlsx")

    password = ask_password() if is_locked(args.pdf) else None

    if args.debug:
        with open_pdf(args.pdf, password) as pdf:
            report = kbank.debug_dump(pdf, args.debug)
        debug_file = args.pdf.with_name("debug.txt")
        debug_file.write_text(report, encoding="utf-8")
        print(f"Wrote {debug_file}")
        return
    for attempt in range(3):
        try:
            stats = convert(args.pdf, out, password)
            break
        except PasswordError:
            if attempt < 2:
                print("Wrong password, try again.")
                password = ask_password()
    else:
        sys.exit("Could not open the PDF.")

    print(f"Done: {out}")
    if stats["format"] == "KBank":
        print(f"  KBank statement, account {stats['account'] or '(not found)'}")
        print(f"  Transactions: {stats['transactions']} on {stats['pages']} pages")
        for label, ok in stats["checks"]:
            status = "OK" if ok else "--" if ok is None else "MISMATCH"
            print(f"  [{status}] {label.split(' / ')[-1]}")  # English part; consoles may not show Thai
    else:
        print(f"  Transactions (date-based lines): {stats['transactions']}")
        print(f"  Table rows detected:             {stats['table_rows']}")
        print("Check both sheets - use whichever matches your statement layout better.")


if __name__ == "__main__":
    main()
