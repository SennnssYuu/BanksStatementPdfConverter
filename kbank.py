"""
Reader for KBank (Kasikornbank) savings account statements
(รายการเดินบัญชีเงินฝากออมทรัพย์).

Main table columns:
    วันที่ | เวลา/วันที่มีผล | รายการ | ถอนเงิน / ฝากเงิน | ยอดคงเหลือ | ช่องทาง | รายละเอียด

Account box at the top of the first page:
    เลขที่อ้างอิง, เลขที่บัญชีเงินฝาก, รอบระหว่างวันที่, สาขาเจ้าของบัญชี,
    ยอดยกไป, รวมถอนเงิน N รายการ, รวมฝากเงิน N รายการ
"""

import re
import statistics
from datetime import datetime

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- Thai text clean-up

# Some Thai PDF fonts store shifted tone marks/vowels in the Private Use Area.
_PUA = {
    0xF700: 0x0E10, 0xF701: 0x0E34, 0xF702: 0x0E35, 0xF703: 0x0E36, 0xF704: 0x0E37,
    0xF705: 0x0E48, 0xF706: 0x0E49, 0xF707: 0x0E4A, 0xF708: 0x0E4B, 0xF709: 0x0E4C,
    0xF70A: 0x0E48, 0xF70B: 0x0E49, 0xF70C: 0x0E4A, 0xF70D: 0x0E4B, 0xF70E: 0x0E4C,
    0xF70F: 0x0E0D, 0xF710: 0x0E31, 0xF711: 0x0E4D, 0xF712: 0x0E47, 0xF713: 0x0E48,
    0xF714: 0x0E49, 0xF715: 0x0E4A, 0xF716: 0x0E4B, 0xF717: 0x0E4C, 0xF718: 0x0E38,
    0xF719: 0x0E39, 0xF71A: 0x0E3A,
}
_THAI_MARKS = "ัิ-ฺ็-๎"


def fix_thai(s):
    s = s.translate(_PUA)
    # Sara am split into nikhahit + sara aa (ํ + า -> ำ), keeping any tone mark before it
    s = re.sub("ํ([่-๋]?)า", lambda m: m.group(1) + "ำ", s)
    # Stray spaces in front of above/below marks
    s = re.sub(f"\\s+(?=[{_THAI_MARKS}])", "", s)
    return s


def _squash(s):
    return re.sub(r"\s+", "", s)


def _join(a, b):
    """Join two text pieces from wrapped lines. Lines wrap at spaces, or right after / or -."""
    if not a:
        return b
    if not b:
        return a
    if a[-1] in "/-":
        return a + b
    return a + " " + b


# ---------------------------------------------------------------- patterns

DATE_RE = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})$")
DATE_ANY_RE = re.compile(r"(?<!\d)\d{1,2}[-/.]\d{1,2}[-/.](?:\d{4}|\d{2})(?!\d)")
TIME_RE = re.compile(r"^\d{1,2}[:.]\d{2}$")
NUM_RE = re.compile(r"^-?[\d,]+\.\d{2}-?$")

COLS = ["date", "time", "desc", "amount", "balance", "channel", "details"]
HEADER_LABELS = {
    "date": "วันที่", "time": "เวลา", "desc": "รายการ", "amount": "ถอนเงิน",
    "balance": "ยอดคงเหลือ", "channel": "ช่องทาง", "details": "รายละเอียด",
}
REQUIRED = {"date", "desc", "amount", "balance"}


def _num(s):
    s = s.replace(",", "").strip()
    neg = s.startswith("-") or s.endswith("-")
    v = float(s.strip("-"))
    return -v if neg else v


def _label(text):
    """Regex for a Thai label that tolerates spaces the PDF may insert between letters."""
    return r"\s*".join(map(re.escape, text.replace(" ", "")))


# ---------------------------------------------------------------- page layout

_MARK_RE = re.compile(f"[{_THAI_MARKS}]")


def _words(page):
    # y_tolerance 5 keeps raised Thai tone marks inside their word
    words = page.extract_words(use_text_flow=True, x_tolerance=2, y_tolerance=5,
                               return_chars=True)
    out = []
    for w in words:
        chars = w.pop("chars")
        w["text"] = fix_thai(w["text"])
        if not w["text"].strip() or not w.get("upright", True):
            continue  # skip sideways text, e.g. the form code printed up the left margin
        # Baseline = where the letters sit. Unlike top/bottom it is the same for Thai and
        # English text on one line, even when they use different fonts.
        base_chars = [c for c in chars if not _MARK_RE.fullmatch(c["text"])] or chars
        w["base"] = statistics.median(page.height - c["matrix"][5] for c in base_chars)
        w["size"] = max(c["size"] for c in chars)
        w["font"] = chars[0].get("fontname", "")
        w["mid"] = (w["top"] + w["bottom"]) / 2
        w["cx"] = (w["x0"] + w["x1"]) / 2
        out.append(w)
    return out


def _find_header(words):
    """Find the table header row. Returns {col: header word} or None."""
    for anchor in [w for w in words if "ยอดคงเหลือ" in _squash(w["text"])]:
        band = [w for w in words if abs(w["mid"] - anchor["mid"]) <= 14]
        found = {}
        for col, label in HEADER_LABELS.items():
            cands = [w for w in band if _squash(w["text"]).startswith(label)]
            if col == "time":  # "เวลา/" or "วันที่มีผล" (the second line of that header)
                cands += [w for w in band if _squash(w["text"]).startswith("วันที่มีผล")]
            if col == "date":
                cands = [w for w in cands if not _squash(w["text"]).startswith("วันที่มีผล")]
            if cands:
                found[col] = min(cands, key=lambda w: w["x0"])
        if REQUIRED <= found.keys():
            return found
    return None


class Layout:
    """Column boundaries for one page."""

    def __init__(self, page, header):
        self.cols = sorted(header, key=lambda c: header[c]["x0"])
        band = list(header.values())
        self.header_bottom = max(w["bottom"] for w in band)
        header_mid = statistics.mean(w["mid"] for w in band)

        vlines = sorted({round(e["x0"], 1) for e in page.edges
                         if e["orientation"] == "v" and e["top"] - 2 <= header_mid <= e["bottom"] + 2})

        self.bounds, self.guessed = [], []
        for a, b in zip(self.cols, self.cols[1:]):
            lo, hi = header[a]["x1"], header[b]["x0"]
            between = [x for x in vlines if lo < x < hi]
            if between:
                self.bounds.append(statistics.median(between))
                self.guessed.append(False)
            else:
                self.bounds.append((lo + hi) / 2)
                self.guessed.append(True)

        # Bottom of the table = lowest vertical grid line at one of the column boundaries
        bottoms = [e["bottom"] for e in page.edges if e["orientation"] == "v"
                   and e["top"] <= header_mid + 2
                   and any(abs(e["x0"] - x) < 2 for x in self.bounds)]
        self.table_bottom = max(bottoms) if bottoms and max(bottoms) > self.header_bottom + 20 \
            else page.height

        # Withdrawal (left) vs deposit (right) inside the amount column
        if "amount" in header:
            h = header["amount"]
            i = self.cols.index("amount")
            left = self.bounds[i - 1] if i > 0 else h["x0"]
            right = self.bounds[i] if i < len(self.bounds) else h["x1"]
            self.amount_split = (left + right) / 2
        else:
            self.amount_split = None

    def col_of(self, w):
        for col, bound in zip(self.cols, self.bounds):
            if w["cx"] < bound:
                return col
        return self.cols[-1]

    def guessed_between(self, a, b):
        if a in self.cols and b in self.cols:
            i = self.cols.index(a)
            if i + 1 < len(self.cols) and self.cols[i + 1] == b:
                return self.guessed[i]
        return False


def _group_lines(words):
    """Group words into visual lines by their baseline."""
    lines = []
    for w in sorted(words, key=lambda w: (w["base"], w["x0"])):
        if lines and abs(lines[-1]["base"] - w["base"]) <= max(2.0, 0.35 * w["size"]):
            lines[-1]["words"].append(w)
        else:
            lines.append({"base": w["base"], "words": [w]})
    for line in lines:
        line["words"].sort(key=lambda w: w["x0"])
        line["base"] = statistics.median(w["base"] for w in line["words"])
        line["size"] = max(w["size"] for w in line["words"])
        line["text"] = " ".join(w["text"] for w in line["words"])
    return lines


def _cell_text(words):
    """Join words in one cell; words that touch are joined without a space."""
    text, prev = "", None
    for w in sorted(words, key=lambda w: w["x0"]):
        gap = w["x0"] - prev["x1"] if prev else 0
        text = text + (" " if prev and gap > 1.5 else "") + w["text"]
        prev = w
    return text.strip()


def _split_cells(line, layout, detail_x):
    cells = {c: [] for c in COLS}
    for w in line["words"]:
        cells[layout.col_of(w)].append(w)

    # Without grid lines the guessed boundaries can be off; fix using what the text looks like.
    if layout.guessed_between("date", "time") or layout.guessed_between("time", "desc"):
        for w in [w for w in cells["date"] if not DATE_RE.match(w["text"])]:
            cells["date"].remove(w)
            cells["time" if TIME_RE.match(w["text"]) else "desc"].append(w)
        for w in [w for w in cells["time"] if not TIME_RE.match(w["text"])]:
            cells["time"].remove(w)
            cells["desc"].append(w)
    if layout.guessed_between("desc", "amount"):
        for w in [w for w in cells["amount"] if not NUM_RE.match(w["text"])]:
            cells["amount"].remove(w)
            cells["desc"].append(w)
    if layout.guessed_between("amount", "balance") or layout.guessed_between("balance", "channel"):
        for w in [w for w in cells["balance"] if not NUM_RE.match(w["text"])]:
            cells["balance"].remove(w)
            cells["channel"].append(w)
        for w in [w for w in cells["channel"] if NUM_RE.match(w["text"])]:
            cells["channel"].remove(w)
            cells["balance"].append(w)
    if layout.guessed_between("channel", "details"):
        right = sorted(cells["channel"] + cells["details"], key=lambda w: w["x0"])
        gaps = [(b["x0"] - a["x1"], i + 1) for i, (a, b) in enumerate(zip(right, right[1:]))]
        big = max(gaps, default=(0, 0))
        if big[0] >= 8:
            cells["channel"], cells["details"] = right[:big[1]], right[big[1]:]
            detail_x.append(right[big[1]]["x0"])
        elif right and detail_x:
            start = statistics.median(detail_x)
            cells["channel"] = [w for w in right if w["x0"] < start - 4]
            cells["details"] = [w for w in right if w["x0"] >= start - 4]

    # Amounts: if two numbers ended up in one column, the right one is the balance
    nums = sorted([w for w in cells["amount"] + cells["balance"] if NUM_RE.match(w["text"])],
                  key=lambda w: w["x0"])
    if len(nums) >= 2:
        cells["amount"], cells["balance"] = nums[:-1], nums[-1:]
    return cells


# ---------------------------------------------------------------- dates

def _parse_date(text, period_year=None):
    m = DATE_RE.match(text)
    if not m:
        return text
    d, mo, y = int(m.group(1)), int(m.group(2)), m.group(3)
    if len(y) == 4:
        year = int(y) - 543 if int(y) > 2400 else int(y)
    else:
        yy = int(y)
        candidates = [2000 + yy, 2000 + yy - 43]  # Gregorian 20yy, or Buddhist 25yy
        year = min(candidates, key=lambda c: abs(c - period_year)) if period_year else 2000 + yy
    try:
        return datetime(year, mo, d)
    except ValueError:
        return text


# ---------------------------------------------------------------- account box

def parse_info(text):
    info = {}

    def grab(key, pattern, conv=None):
        m = re.search(pattern, text, re.MULTILINE)
        if m:
            vals = [g.strip() for g in m.groups()]
            info[key] = conv(*vals) if conv else vals[0]

    grab("name", _label("ชื่อบัญชี") + r"\s*(.+?)\s*$")
    grab("account", _label("เลขที่บัญชีเงินฝาก") + r"\s*([\dXx][\dXx\- ]{5,}[\dXx])")
    grab("reference", _label("เลขที่อ้างอิง") + r"\s*([A-Za-z0-9]{6,})")
    grab("period", _label("รอบระหว่างวันที่") +
         r"\s*(\d{1,2}/\d{1,2}/\d{4})\s*[-–]\s*(\d{1,2}/\d{1,2}/\d{4})",
         lambda a, b: f"{a} - {b}")
    grab("branch", _label("สาขาเจ้าของบัญชี") + r"\s*(.+?)\s*$")
    grab("closing", _label("ยอดยกไป") + r"\s*(-?[\d,]+\.\d{2})", _num)
    grab("withdrawals", _label("รวมถอนเงิน") + r"\s*([\d,]+)\s*" + _label("รายการ") +
         r"\s*([\d,]+\.\d{2})", lambda n, a: (int(n.replace(",", "")), _num(a)))
    grab("deposits", _label("รวมฝากเงิน") + r"\s*([\d,]+)\s*" + _label("รายการ") +
         r"\s*([\d,]+\.\d{2})", lambda n, a: (int(n.replace(",", "")), _num(a)))
    if "account" in info:
        info["account"] = info["account"].replace(" ", "")
    # Name/branch can run into the account box printed on the same line; cut at the next label
    other = "|".join(_label(x) for x in ("เลขที่", "รอบระหว่าง", "สาขาเจ้าของ", "ยอดยก",
                                         "รวมถอน", "รวมฝาก", "หน้าที่"))
    for key in ("name", "branch"):
        if key in info:
            info[key] = re.split(other, info[key])[0].strip()
    return info


# ---------------------------------------------------------------- diagnostics

def debug_dump(pdf, needle, around=3):
    """Describe how the lines near `needle` (e.g. a date like 04-09-26) are read."""
    out, layout = [], None
    for page_no, page in enumerate(pdf.pages, 1):
        words = _words(page)
        header = _find_header(words)
        if header:
            layout = Layout(page, header)
            out.append(f"--- page {page_no}: columns {layout.cols}, boundaries "
                       f"{[round(b, 1) for b in layout.bounds]}, guessed {layout.guessed}, "
                       f"table bottom {round(layout.table_bottom, 1)}")
        if layout is None:
            continue
        lines = _group_lines(words)
        hits = [i for i, l in enumerate(lines) if needle in l["text"]]
        shown = set()
        for i in hits:
            for j in range(max(0, i - around), min(len(lines), i + around + 1)):
                if j in shown:
                    continue
                shown.add(j)
                line = lines[j]
                cells = _split_cells(line, layout, [])
                col = {id(w): c for c, ws in cells.items() for w in ws}
                out.append(f"line base={line['base']:.1f} size={line['size']:.1f}")
                for w in line["words"]:
                    out.append(f"    {col.get(id(w), '?'):8} x={w['x0']:6.1f}-{w['x1']:6.1f} "
                               f"top={w['top']:6.1f} base={w['base']:6.1f} "
                               f"font={w['font'][-20:]:20} {w['text']}")
    return "\n".join(out) if out else "Table header not found."


# ---------------------------------------------------------------- main parse

def parse(pdf, progress=None):
    """
    Returns a dict with info / transactions / raw lines,
    or None if the PDF doesn't look like a KBank statement.
    """
    raw, rows, info_text, skipped = [], [], [], []
    layout, detail_x = None, []
    total = len(pdf.pages)

    for page_no, page in enumerate(pdf.pages, 1):
        text = fix_thai(page.extract_text() or "")
        raw += [[page_no, line] for line in text.splitlines()]
        if page_no <= 2:
            info_text.append(text)

        words = _words(page)
        header = _find_header(words)
        if header:
            layout = Layout(page, header)
            body = [w for w in words if w["top"] >= layout.header_bottom - 1
                    and w["bottom"] <= layout.table_bottom + 1]
        elif layout is not None:  # page without its own header: reuse previous columns
            body = words
        else:
            if progress:
                progress(page_no, total)
            continue

        cur, last_base, page_has_rows = None, None, False
        for line in _group_lines(body):
            cells = _split_cells(line, layout, detail_x)
            date_word = next((w for w in cells["date"] if DATE_RE.match(w["text"])), None)
            if date_word:  # anything else that fell in the date column belongs to the description
                cells["desc"] = [w for w in cells["date"] if w is not date_word] + cells["desc"]
                cells["date"] = [date_word]
            txt = {c: _cell_text(cells[c]) for c in COLS}
            amount_word = next((w for w in cells["amount"] if NUM_RE.match(w["text"])), None)
            fields = {
                "amount": _num(amount_word["text"]) if amount_word else None,
                "amount_left": (amount_word["cx"] < layout.amount_split)
                if amount_word and layout.amount_split else None,
                "balance": _num(txt["balance"]) if NUM_RE.match(txt["balance"]) else None,
            }
            close = cur is not None and line["base"] - last_base <= 2.5 * line["size"]

            if date_word:
                cur = {"page": page_no, "date": txt["date"], "time": txt["time"],
                       "desc": txt["desc"], "channel": txt["channel"], "details": txt["details"],
                       **fields}
                rows.append(cur)
                page_has_rows = True
            elif close and cur["balance"] is None and fields["balance"] is not None:
                # Date line and amount line came out as two lines: put them back together
                for c in ("time", "desc", "channel", "details"):
                    cur[c] = _join(cur[c], txt[c])
                cur.update(fields)
            elif close and not (txt["time"] or txt["amount"] or txt["balance"]):
                # Wrapped text belonging to the transaction above
                for c in ("desc", "channel", "details"):
                    cur[c] = _join(cur[c], txt[c])
            else:
                cur = None
                looks_like_data = DATE_ANY_RE.search(line["text"]) or any(
                    NUM_RE.match(w["text"]) for w in line["words"])
                is_total = re.search(_label("ยอดยก") + "|" + _label("รวม"), line["text"])
                if page_has_rows and looks_like_data and not is_total:
                    skipped.append({"page": page_no, "text": line["text"]})
                continue
            last_base = line["base"]

        if progress:
            progress(page_no, total)

    if layout is None:
        return None

    info = parse_info("\n".join(info_text))
    period_year = None
    if "period" in info:
        period_year = int(info["period"].split(" - ")[0][-4:])
    return {"info": info, "transactions": _finish(rows, period_year), "raw": raw,
            "skipped": skipped}


def _finish(rows, period_year):
    """Split amounts into withdrawal/deposit using the balance, and check every balance."""
    out, prev = [], None
    for r in rows:
        desc = _squash(r["desc"])
        is_bf = desc.startswith("ยอดยกมา") and r["amount"] is None
        is_cf = desc.startswith("ยอดยกไป") and r["amount"] is None
        if is_cf:
            continue
        if is_bf:
            if out:  # brought-forward line repeated on a later page
                continue
            r.update(withdrawal=None, deposit=None, check=None, opening=True)
        else:
            r["opening"] = False
            r["withdrawal"] = r["deposit"] = None
            amt = r["amount"]
            if amt is not None:
                diff = None if prev is None or r["balance"] is None else round(r["balance"] - prev, 2)
                if diff is not None and abs(abs(diff) - abs(amt)) < 0.005:
                    is_withdrawal = diff < 0
                elif r["amount_left"] is not None:
                    is_withdrawal = r["amount_left"]
                else:
                    is_withdrawal = amt < 0
                if is_withdrawal:
                    r["withdrawal"] = abs(amt)
                else:
                    r["deposit"] = abs(amt)
            if prev is None or r["balance"] is None or amt is None:
                r["check"] = None
            else:
                expected = prev - (r["withdrawal"] or 0) + (r["deposit"] or 0)
                r["check"] = abs(expected - r["balance"]) < 0.005
        r["date"] = _parse_date(r["date"], period_year)
        if r["balance"] is not None:
            prev = r["balance"]
        out.append(r)
    return out


def summarize(result):
    """Compare the statement's own totals with what was extracted."""
    info, txs = result["info"], result["transactions"]
    w = [t["withdrawal"] for t in txs if t["withdrawal"] is not None]
    d = [t["deposit"] for t in txs if t["deposit"] is not None]
    last_balance = next((t["balance"] for t in reversed(txs) if t["balance"] is not None), None)
    bad_rows = sum(1 for t in txs if t["check"] is False)

    checks = []

    def add(label, stated, found, money):
        if stated is None:
            ok = None
        else:
            ok = abs(stated - found) < 0.005 if money else stated == found
        checks.append((label, stated, found, money, ok))

    add("จำนวนรายการถอน / Withdrawals (count)", info.get("withdrawals", (None,))[0], len(w), False)
    add("ยอดรวมถอนเงิน / Total withdrawn", info.get("withdrawals", (None, None))[1], round(sum(w), 2), True)
    add("จำนวนรายการฝาก / Deposits (count)", info.get("deposits", (None,))[0], len(d), False)
    add("ยอดรวมฝากเงิน / Total deposited", info.get("deposits", (None, None))[1], round(sum(d), 2), True)
    add("ยอดยกไป / Closing balance", info.get("closing"), last_balance, True)
    checks.append(("ยอดคงเหลือต่อเนื่องทุกบรรทัด / Running balance (rows wrong)",
                   0, bad_rows, False, bad_rows == 0))
    skipped = len(result.get("skipped", []))
    checks.append(("บรรทัดที่อ่านไม่ได้ / Unreadable lines in table",
                   0, skipped, False, skipped == 0))
    return checks


# ---------------------------------------------------------------- Excel output

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
BAD_FILL = PatternFill("solid", fgColor="FDE2E1")
OK_FONT = Font(color="1E7B45", bold=True)
BAD_FONT = Font(color="B42318", bold=True)
THIN = Side(style="thin", color="D0D7E2")
MONEY = "#,##0.00"


def _header_row(ws, row, labels):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN)


def write_workbook(wb, result):
    info, txs = result["info"], result["transactions"]
    account = info.get("account", "")

    # --- Transactions
    ws = wb.active
    ws.title = "Transactions"
    cols = [("เลขที่บัญชี\nAccount No.", 16), ("วันที่\nDate", 12), ("เวลา\nTime", 8),
            ("รายการ\nDescription", 32), ("ถอนเงิน\nWithdrawal", 15), ("ฝากเงิน\nDeposit", 15),
            ("ยอดคงเหลือ\nBalance", 17), ("ช่องทาง\nChannel", 26), ("รายละเอียด\nDetails", 44),
            ("หน้า\nPage", 7), ("ตรวจยอด\nCheck", 10)]
    _header_row(ws, 1, [c[0] for c in cols])
    ws.row_dimensions[1].height = 32
    for i, (_, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for r, t in enumerate(txs, 2):
        check = "" if t["check"] is None else ("✓" if t["check"] else "✗ ยอดไม่ตรง")
        values = [account, t["date"], t["time"], t["desc"], t["withdrawal"], t["deposit"],
                  t["balance"], t["channel"], t["details"], t["page"], check]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=col in (4, 8, 9))
        ws.cell(row=r, column=2).number_format = "dd/mm/yyyy"
        for col in (5, 6, 7):
            ws.cell(row=r, column=col).number_format = MONEY
        ws.cell(row=r, column=5).font = Font(color="B42318")
        ws.cell(row=r, column=6).font = Font(color="1E7B45")
        ws.cell(row=r, column=11).alignment = Alignment(horizontal="center", vertical="top")
        if t["check"] is False:
            for col in range(1, len(cols) + 1):
                ws.cell(row=r, column=col).fill = BAD_FILL

    last = len(txs) + 1
    total_row = last + 1
    ws.cell(row=total_row, column=4, value="รวม / Total").font = Font(bold=True)
    for col in (5, 6):
        letter = get_column_letter(col)
        c = ws.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{last})")
        c.number_format, c.font = MONEY, Font(bold=True)
        c.border = Border(top=Side(style="thin"), bottom=Side(style="double"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    # --- Account info + checks
    wi = wb.create_sheet("Account Info")
    wi.column_dimensions["A"].width = 46
    wi.column_dimensions["B"].width = 34
    wi.column_dimensions["C"].width = 20
    wi.column_dimensions["D"].width = 12
    _header_row(wi, 1, ["ข้อมูลบัญชี / Account", "", "", ""])
    wi.merge_cells("A1:D1")
    opening = next((t["balance"] for t in txs if t.get("opening")), None)
    fields = [("ชื่อบัญชี / Account name", info.get("name")),
              ("เลขที่บัญชีเงินฝาก / Account no.", info.get("account")),
              ("เลขที่อ้างอิง / Reference no.", info.get("reference")),
              ("รอบระหว่างวันที่ / Period", info.get("period")),
              ("สาขาเจ้าของบัญชี / Branch", info.get("branch")),
              ("ยอดยกมา / Opening balance", opening),
              ("ยอดยกไป / Closing balance", info.get("closing"))]
    for i, (label, value) in enumerate(fields, 2):
        wi.cell(row=i, column=1, value=label).font = Font(bold=True)
        c = wi.cell(row=i, column=2, value=value if value is not None else "—")
        if isinstance(value, float):
            c.number_format = MONEY
        c.alignment = Alignment(horizontal="left")

    start = len(fields) + 3
    _header_row(wi, start, ["ตรวจสอบ / Check", "ตามใบแจ้งยอด / Statement",
                            "จากไฟล์นี้ / Extracted", "ผล / Result"])
    wi.row_dimensions[start].height = 30
    for i, (label, stated, found, money, ok) in enumerate(summarize(result), start + 1):
        wi.cell(row=i, column=1, value=label)
        for col, v in ((2, stated), (3, found)):
            c = wi.cell(row=i, column=col, value=v if v is not None else "—")
            if money and isinstance(v, (int, float)):
                c.number_format = MONEY
            c.alignment = Alignment(horizontal="right")
        res = wi.cell(row=i, column=4, value="—" if ok is None else ("✓ ตรง" if ok else "✗ ไม่ตรง"))
        res.font = OK_FONT if ok else (BAD_FONT if ok is False else Font(color="5E6B7D"))
        res.alignment = Alignment(horizontal="center")

    # --- Lines inside the table that couldn't be read as a transaction
    if result.get("skipped"):
        wsk = wb.create_sheet("Skipped lines", 1)
        wsk.sheet_properties.tabColor = "B42318"
        _header_row(wsk, 1, ["หน้า / Page", "บรรทัดที่อ่านไม่ได้ / Line that couldn't be read"])
        wsk.column_dimensions["A"].width = 10
        wsk.column_dimensions["B"].width = 120
        for s in result["skipped"]:
            wsk.append([s["page"], s["text"]])

    # --- Raw text (for checking)
    wr = wb.create_sheet("Raw text")
    _header_row(wr, 1, ["หน้า / Page", "ข้อความ / Line"])
    wr.column_dimensions["A"].width = 10
    wr.column_dimensions["B"].width = 120
    for row in result["raw"]:
        wr.append(row)
